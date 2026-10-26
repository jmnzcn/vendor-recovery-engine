"""Client-facing outputs: Excel workbook, PDF summary and dispute-ready CSV per vendor,
plus a portfolio summary. Written in English for the client; amounts in USD unless noted."""
from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from vre.policy import Category, Status
from vre.triage import queue_summary

CATEGORY_LABEL = {
    Category.SHORTAGE: "Shortage claims", Category.PRICE_CLAIM: "Price claims",
    Category.CHARGEBACK_LATE_SHIPMENT: "Late-shipment chargebacks", Category.CHARGEBACK_ASN: "ASN chargebacks",
    Category.DUPLICATE_DEDUCTION: "Duplicate deductions", Category.UNDERPAYMENT: "Underpayments",
    Category.UNPAID_INVOICE: "Unpaid invoices", Category.FX_VARIANCE: "Exchange-rate variances",
}
STATUS_LABEL = {
    Status.CLAIMABLE: "Claimable", Status.NOT_CLAIMABLE: "Valid deduction", Status.INSUFFICIENT_EVIDENCE:
    "Insufficient evidence", Status.NEEDS_REVIEW: "Needs review", Status.ALREADY_RECOVERED: "Already recovered",
    Status.EXPIRED: "Past dispute window",
}
NAVY, LIGHT = "1F3A5F", "E8EEF5"


def usd(cents: float) -> str:
    return f"${cents / 100:,.2f}"


def summarize(findings: pd.DataFrame) -> pd.DataFrame:
    g = findings.groupby("category")
    out = pd.DataFrame({
        "findings": g.size(),
        "claimable": g["status"].apply(lambda s: int((s == Status.CLAIMABLE).sum())),
        "recoverable_usd": g["recoverable_usd_cents"].sum() / 100,
        "expired_usd": g["expired_usd_cents"].sum() / 100,
    })
    return out.reindex([c.value for c in Category]).dropna(how="all").fillna(0)


def dispute_csv(findings: pd.DataFrame, path: Path) -> pd.DataFrame:
    c = findings[findings["status"] == Status.CLAIMABLE].sort_values("priority_rank")
    out = pd.DataFrame({
        "priority": c["priority_rank"], "urgent": c["urgent"], "days_left": c["days_left"],
        "reference_type": c["ref_type"], "reference": c["ref"], "category": c["category"],
        "dispute_amount": c["amount_cents"].where(c["ref_type"] == "deduction", c["recoverable_usd_cents"]) / 100,
        "currency": c["currency"].where(c["ref_type"] == "deduction", "USD"),
        "recoverable_usd": c["recoverable_usd_cents"] / 100,
        "expected_value_usd": c["expected_value_usd_cents"] / 100,
        "dispute_by": c["dispute_deadline"].dt.strftime("%Y-%m-%d"),
        "reason": c["reason"], "evidence": c["evidence"],
        "source": c["source_file"] + ":" + c["source_row"].astype("Int64").astype(str),
    })
    out.to_csv(path, index=False)
    return out


def _sheet(ws, df: pd.DataFrame, money_cols: set[str], widths: dict[str, int] | None = None) -> None:
    ws.append(list(df.columns))
    for row in df.itertuples(index=False):
        ws.append([None if pd.isna(v) else v for v in row])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for j, col in enumerate(df.columns, start=1):
        letter_ = get_column_letter(j)
        ws.column_dimensions[letter_].width = (widths or {}).get(col, min(60, max(12, len(col) + 2)))
        if col in money_cols:
            for cell in ws[letter_][1:]:
                cell.number_format = '"$"#,##0.00'
    ws.freeze_panes = "A2"
    if len(df):
        ws.auto_filter.ref = ws.dimensions


def excel(findings: pd.DataFrame, vendor: dict, quality: pd.DataFrame, path: Path, as_of: str) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    s = summarize(findings)
    ws.append([f"Recovery report: {vendor['vendor_name']} ({vendor['vendor_code']})"])
    ws["A1"].font = Font(bold=True, size=14, color=NAVY)
    ws.append([f"Data as of {as_of}. Amounts in USD."])
    ws.append([])
    ws.append(["Recoverable now", findings["recoverable_usd_cents"].sum() / 100])
    ws.append(["Lost to the dispute window", findings["expired_usd_cents"].sum() / 100])
    ws.append(["Claims ready to file", int((findings["status"] == Status.CLAIMABLE).sum())])
    q = queue_summary(findings)
    ws.append(["Expiring in the next 30 days", q["urgent_usd"]])
    ws.append(["Expected value of all claims", q["expected_value_usd"]])
    for r in (4, 5, 7, 8):
        ws.cell(r, 2).number_format = '"$"#,##0.00'
        ws.cell(r, 1).font = Font(bold=True)
    ws.append([])
    start = ws.max_row + 1
    ws.append(["Category", "Findings", "Claimable", "Recoverable (USD)", "Expired (USD)"])
    for cat, r in s.iterrows():
        ws.append([CATEGORY_LABEL.get(Category(cat), cat), int(r["findings"]), int(r["claimable"]),
                   r["recoverable_usd"], r["expired_usd"]])
    for cell in ws[start]:
        cell.font, cell.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor=NAVY)
    for row in ws.iter_rows(min_row=start + 1, min_col=4, max_col=5):
        for cell in row:
            cell.number_format = '"$"#,##0.00'
    for col, w in zip("ABCDE", (34, 12, 12, 20, 18), strict=True):
        ws.column_dimensions[col].width = w

    cols = ["priority_rank", "urgent", "days_left", "ref_type", "ref", "category", "status", "recoverable_usd",
            "expected_value_usd", "amount", "currency", "dispute_deadline", "match_tier", "reason", "evidence",
            "source_file", "source_row"]
    f = findings.assign(recoverable_usd=findings["recoverable_usd_cents"] / 100,
                        expected_value_usd=findings["expected_value_usd_cents"] / 100,
                        amount=findings["amount_cents"] / 100,
                        dispute_deadline=findings["dispute_deadline"].dt.strftime("%Y-%m-%d"))[cols]
    wide = {"reason": 60, "evidence": 70, "source_file": 26}
    for title, mask in [
        ("Work queue", f["status"] == Status.CLAIMABLE),
        ("Needs review", f["status"] == Status.NEEDS_REVIEW),
        ("Insufficient evidence", f["status"] == Status.INSUFFICIENT_EVIDENCE),
        ("All findings", pd.Series(True, index=f.index)),
    ]:
        sub = f[mask].sort_values("priority_rank") if title == "Work queue" else f[mask]
        _sheet(wb.create_sheet(title), sub, {"recoverable_usd", "expected_value_usd"}, wide)
    _sheet(wb.create_sheet("Data quality"), quality, set())
    wb.save(path)


def pdf(findings: pd.DataFrame, vendor: dict, quality: pd.DataFrame, path: Path, as_of: str) -> None:
    ss = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=ss["Title"], alignment=0, fontSize=18, textColor=colors.HexColor(f"#{NAVY}"))
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], textColor=colors.HexColor(f"#{NAVY}"), spaceBefore=10)
    body = ParagraphStyle("b", parent=ss["BodyText"], fontSize=9.5, leading=13)
    small = ParagraphStyle("s", parent=body, fontSize=7.5, leading=10, textColor=colors.HexColor("#555555"))
    cell = ParagraphStyle("c", parent=body, fontSize=8, leading=10)

    def table(rows: list[list], widths: list[float], right: tuple[int, ...] = ()) -> Table:
        t = Table(rows, colWidths=widths, repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{NAVY}")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("FONTSIZE", (0, 0), (-1, -1), 8.5),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor(f"#{LIGHT}")]),
            *[("ALIGN", (c, 0), (c, -1), "RIGHT") for c in right], ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4), ("TOPPADDING", (0, 0), (-1, -1), 4),
        ]))
        return t

    rec = findings["recoverable_usd_cents"].sum()
    exp = findings["expired_usd_cents"].sum()
    n_claims = int((findings["status"] == Status.CLAIMABLE).sum())
    q = queue_summary(findings)
    urgent_line = (f" <b>{usd(q['urgent_usd'] * 100)}</b> of it expires in the next 30 days "
                   f"({q['urgent_claims']:,} claims): file those first." if q["urgent_claims"] else "")
    story = [
        Paragraph(f"Recovery report: {vendor['vendor_name']}", h1),
        Paragraph(f"Vendor code {vendor['vendor_code']} · data as of {as_of} · amounts in USD", small),
        Spacer(1, 10),
        Paragraph(f"We reviewed <b>{len(findings):,}</b> deductions and payments. <b>{usd(rec)}</b> can be "
                  f"disputed now across <b>{n_claims:,}</b> claims, each backed by your own shipping, invoice or "
                  f"payment records.{urgent_line} Another <b>{usd(exp)}</b> was recoverable but is past the "
                  f"dispute window.", body),
        Paragraph("Recoverable by category", h2),
    ]
    s = summarize(findings)
    rows = [["Category", "Findings", "Claimable", "Recoverable", "Expired"]] + [
        [CATEGORY_LABEL.get(Category(c), c), f"{int(r['findings']):,}", f"{int(r['claimable']):,}",
         usd(r["recoverable_usd"] * 100), usd(r["expired_usd"] * 100)] for c, r in s.iterrows()]
    rows.append(["Total", f"{int(s['findings'].sum()):,}", f"{int(s['claimable'].sum()):,}", usd(rec), usd(exp)])
    story.append(table(rows, [2.6 * inch, 0.9 * inch, 0.9 * inch, 1.3 * inch, 1.2 * inch], right=(1, 2, 3, 4)))

    top = findings[findings["status"] == Status.CLAIMABLE].nsmallest(10, "priority_rank")
    if len(top):
        story += [Paragraph("File these first", h2), Paragraph(
            "Ranked by deadline (30 days or less first), then by expected value: amount times an assumed win "
            "rate for the kind of evidence.", small), Spacer(1, 4)]
        rows = [["Reference", "Category", "Amount", "Days left", "Why"]] + [
            [r.ref, CATEGORY_LABEL.get(Category(r.category), r.category), usd(r.recoverable_usd_cents),
             f"{int(r.days_left)}{' !' if r.urgent else ''}", Paragraph(r.reason, cell)] for r in top.itertuples()]
        story.append(table(rows, [1.15 * inch, 1.45 * inch, 0.85 * inch, 0.7 * inch, 2.85 * inch], right=(2, 3)))

    st = findings["status"].value_counts()
    story += [Paragraph("What we did not claim, and why", h2), Paragraph(
        f"<b>{st.get(Status.NOT_CLAIMABLE, 0):,}</b> deductions are valid: your records confirm them. "
        f"<b>{st.get(Status.INSUFFICIENT_EVIDENCE, 0):,}</b> cannot be proven either way (for example, no carrier "
        f"proof of delivery). <b>{st.get(Status.NEEDS_REVIEW, 0):,}</b> quote a reference that matches more than "
        f"one of your records, so they were not matched automatically. "
        f"<b>{st.get(Status.ALREADY_RECOVERED, 0):,}</b> were already reversed by Amazon.", body)]
    q = quality.agg({"rows_read": "sum", "rows_quarantined": "sum", "duplicates_removed": "sum"})
    story += [Paragraph("Data quality", h2), Paragraph(
        f"{int(q['rows_read']):,} rows read from your exports. {int(q['rows_quarantined']):,} rows were set aside "
        f"(summary lines, repeated headers or rows missing required fields) and "
        f"{int(q['duplicates_removed']):,} duplicated export rows were counted once. Every claim lists the file "
        f"and row it comes from in the accompanying workbook.", body),
        Spacer(1, 14),
        Paragraph("Dispute windows and tolerances are configurable assumptions; confirm current Vendor Central "
                  "terms before filing. Exchange-rate checks use ECB reference rates.", small)]
    SimpleDocTemplate(str(path), pagesize=letter, leftMargin=0.7 * inch, rightMargin=0.7 * inch,
                      topMargin=0.7 * inch, bottomMargin=0.7 * inch,
                      title=f"Recovery report {vendor['vendor_code']}").build(story)


def write_all(findings: pd.DataFrame, vendors: pd.DataFrame, stats: pd.DataFrame, out: Path, as_of: str,
              pdf_vendors: list[str] | None = None) -> dict:
    """Per-vendor reports plus portfolio files. PDFs for `pdf_vendors` (default: all)."""
    out.mkdir(parents=True, exist_ok=True)
    (out / "vendors").mkdir(exist_ok=True)
    findings.to_parquet(out / "findings.parquet", index=False)
    dispute_csv(findings, out / "claims_all_vendors.csv")
    names = vendors.set_index("vendor_code")["vendor_name"].to_dict() if len(vendors) else {}
    for code, f in findings.groupby("vendor_code"):
        v = {"vendor_code": code, "vendor_name": names.get(code, code)}
        q = stats[stats["vendor_code"] == code][["table", "file", "rows_read", "rows_kept", "rows_quarantined",
                                                 "duplicates_removed"]]
        vdir = out / "vendors" / code
        vdir.mkdir(exist_ok=True)
        excel(f, v, q, vdir / f"recovery_{code}.xlsx", as_of)
        dispute_csv(f, vdir / f"claims_{code}.csv")
        if pdf_vendors is None or code in pdf_vendors:
            pdf(f, v, q, vdir / f"recovery_{code}.pdf", as_of)

    by_vendor = findings.groupby("vendor_code").agg(
        findings=("ref", "size"), claims=("status", lambda s: int((s == Status.CLAIMABLE).sum())),
        recoverable_usd=("recoverable_usd_cents", lambda s: s.sum() / 100),
        expired_usd=("expired_usd_cents", lambda s: s.sum() / 100)).sort_values("recoverable_usd", ascending=False)
    by_vendor.insert(0, "vendor_name", by_vendor.index.map(lambda c: names.get(c, c)))
    by_vendor.to_csv(out / "portfolio_by_vendor.csv")
    summary = {
        "as_of": as_of,
        "vendors": int(findings["vendor_code"].nunique()),
        "findings": len(findings),
        "claims": int((findings["status"] == Status.CLAIMABLE).sum()),
        "recoverable_usd": round(findings["recoverable_usd_cents"].sum() / 100, 2),
        "expired_usd": round(findings["expired_usd_cents"].sum() / 100, 2),
        **queue_summary(findings),
        "by_category": summarize(findings).reset_index().to_dict("records"),
        "by_status": findings["status"].value_counts().to_dict(),
    }
    (out / "summary.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")
    return summary
